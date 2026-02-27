"""
CYBER FORGE SCAN - Convertisseur DOCX / TXT vers PDF
Unicode + Numérotation des pages 
"""

import os
import sys
import subprocess
import re
from typing import Optional, List


# ================= VÉRIFICATION ET INSTALLATION DES DÉPENDANCES =================

def check_and_install_dependencies():
    """Vérifie et installe les dépendances nécessaires"""
    dependencies = {
        "fpdf": "fpdf2",
        "docx": "python-docx"
    }

    print("\n📦 Vérification des dépendances...")
    
    for module, package in dependencies.items():
        try:
            __import__(module)
            print(f"✅ {package} déjà installé")
        except ImportError:
            print(f"📦 Installation de {package}...")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", package, "--quiet"])
                print(f"✅ {package} installé avec succès")
            except subprocess.CalledProcessError:
                print(f"❌ Erreur lors de l'installation de {package}")
                return False
    
    return True


# ================= NETTOYAGE DU TEXTE =================

def clean_text(text: str) -> str:
    """
    Nettoie le texte en remplaçant les emojis et caractères non supportés
    
    Args:
        text: Texte à nettoyer
    
    Returns:
        Texte nettoyé
    """
    # Dictionnaire de remplacement des emojis courants
    emoji_map = {
        '🔑': '[CLE]', '✅': '[OK]', '❌': '[X]', '⚠️': '[!]',
        '📁': '[DOSSIER]', '📄': '[FICHIER]', '💡': '[IDEE]',
        '🚀': '[FUSEE]', '📊': '[STATS]', '⏳': '[ATTENTE]',
        '🔐': '[SECURITE]', '🔒': '[VERROU]', '🔓': '[OUVERT]',
        '💻': '[ORDI]', '📱': '[MOBILE]', '🌐': '[WEB]',
        '📧': '[EMAIL]', '📞': '[TEL]', '🏠': '[MAISON]',
        '⭐': '[STAR]', '✨': '[SPARKLE]', '🎯': '[CIBLE]',
        '📝': '[NOTE]', '📖': '[LIVRE]', '🎓': '[DIPLOME]',
        '🔍': '[RECHERCHE]', '🛡️': '[BOUCLIER]', '⚡': '[ECLAIR]'
    }

    # Remplacer les emojis
    for emoji, replacement in emoji_map.items():
        text = text.replace(emoji, replacement)

    # Nettoyer les caractères non-ASCII (garder les caractères accentués français)
    # Conserver: espace, tab, retour ligne, ASCII imprimable, et caractères latins étendus
    text = re.sub(r'[^\x09\x0A\x0D\x20-\x7E\u00A0-\u024F]', '', text)
    
    return text.strip()


# ================= LECTURE DES FICHIERS =================

def lire_docx(path: str) -> Optional[List[str]]:
    """
    Lit un fichier DOCX et retourne les lignes de texte
    
    Args:
        path: Chemin du fichier DOCX
    
    Returns:
        Liste de lignes ou None si erreur
    """
    try:
        from docx import Document
        doc = Document(path)
        lines = [clean_text(p.text) for p in doc.paragraphs if p.text.strip()]
        return lines
    except Exception as e:
        print(f"❌ Erreur lors de la lecture du DOCX: {e}")
        return None


def lire_txt(path: str) -> Optional[List[str]]:
    """
    Lit un fichier TXT avec détection automatique de l'encodage
    
    Args:
        path: Chemin du fichier TXT
    
    Returns:
        Liste de lignes ou None si erreur
    """
    encodages = ["utf-8", "cp1252", "latin-1", "iso-8859-1"]
    
    for encodage in encodages:
        try:
            with open(path, "r", encoding=encodage) as f:
                lines = [clean_text(line) for line in f if line.strip()]
                print(f"✅ Fichier lu avec l'encodage: {encodage}")
                return lines
        except UnicodeDecodeError:
            continue
        except Exception as e:
            print(f"❌ Erreur lors de la lecture: {e}")
            return None
    
    print("❌ Impossible de détecter l'encodage du fichier")
    return None


# ================= CLASSE PDF PERSONNALISÉE =================

try:
    from fpdf import FPDF

    class CustomPDF(FPDF):
        """Classe PDF personnalisée avec en-tête et pied de page"""
        
        def __init__(self, titre: str = "Document"):
            super().__init__()
            self.titre_document = titre
        
        def header(self):
            """En-tête personnalisé"""
            try:
                self.set_font("DejaVu", size=10)
                self.set_text_color(100, 100, 100)
                self.cell(0, 10, self.titre_document, align="C", ln=True)
                self.ln(2)
            except:
                pass
        
        def footer(self):
            """Pied de page avec numéro"""
            self.set_y(-15)
            try:
                self.set_font("DejaVu", size=9)
                self.set_text_color(100, 100, 100)
                self.cell(0, 10, f"Page {self.page_no()}", align="C")
            except:
                pass

except ImportError:
    print("⚠️  Module fpdf2 non disponible")
    CustomPDF = None


# ================= CONVERSION VERS PDF =================

def convert_to_pdf(source_path: str, output_name: str, output_dir: Optional[str] = None) -> bool:
    """
    Convertit un fichier DOCX ou TXT en PDF
    
    Args:
        source_path: Chemin du fichier source
        output_name: Nom du fichier de sortie (sans extension)
        output_dir: Dossier de sortie (optionnel)
    
    Returns:
        True si succès, False sinon
    """
    # Vérifier que les dépendances sont installées
    if CustomPDF is None:
        print("❌ Dépendances manquantes. Exécutez check_and_install_dependencies() d'abord.")
        return False
    
    # Vérifier l'existence du fichier
    if not os.path.exists(source_path):
        print(f"❌ Fichier introuvable: {source_path}")
        return False

    # Vérifier l'extension
    ext = os.path.splitext(source_path)[1].lower()
    if ext not in (".docx", ".txt"):
        print(f"❌ Format non supporté: {ext}")
        print("   Formats acceptés: .docx, .txt")
        return False

    # Lire le fichier
    print(f"\n📖 Lecture du fichier: {os.path.basename(source_path)}")
    if ext == ".docx":
        lines = lire_docx(source_path)
    else:
        lines = lire_txt(source_path)
    
    if not lines:
        print("❌ Impossible de lire le fichier ou fichier vide")
        return False
    
    print(f"✅ {len(lines)} lignes lues")

    # Déterminer le dossier de sortie
    if output_dir is None:
        output_dir = os.path.dirname(source_path) or "."
    
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, output_name + ".pdf")

    # Vérifier la présence de la police
    font_path = os.path.join(os.path.dirname(__file__), "DejaVuSans.ttf")
    if not os.path.exists(font_path):
        print("⚠️  Police DejaVuSans.ttf non trouvée")
        print("   Téléchargement recommandé depuis: https://dejavu-fonts.github.io/")
        print("   Tentative de conversion sans police personnalisée...")
        
        # Utiliser Arial comme fallback
        try:
            pdf = CustomPDF(titre=os.path.basename(source_path))
            pdf.set_auto_page_break(auto=True, margin=20)
            pdf.add_page()
            pdf.set_font("Arial", size=11)
        except:
            print("❌ Impossible de créer le PDF")
            return False
    else:
        # Créer le PDF avec la police personnalisée
        pdf = CustomPDF(titre=os.path.basename(source_path))
        pdf.set_auto_page_break(auto=True, margin=20)
        pdf.add_page()
        
        try:
            pdf.add_font("DejaVu", "", font_path, uni=True)
            pdf.set_font("DejaVu", size=11)
        except Exception as e:
            print(f"⚠️  Erreur lors du chargement de la police: {e}")
            pdf.set_font("Arial", size=11)

    # Ajouter le titre
    try:
        pdf.set_font("DejaVu", size=16)
    except:
        pdf.set_font("Arial", size=16)
    
    pdf.multi_cell(0, 10, clean_text(os.path.basename(source_path)), align="C")
    pdf.ln(5)

    # Restaurer la taille normale
    try:
        pdf.set_font("DejaVu", size=11)
    except:
        pdf.set_font("Arial", size=11)

    # Ajouter le contenu
    print("📝 Génération du PDF...")
    for i, line in enumerate(lines, 1):
        try:
            if line:  # Ignorer les lignes vides
                pdf.multi_cell(0, 7, line)
                pdf.ln(1)
        except Exception as e:
            print(f"⚠️  Erreur à la ligne {i}: {e}")
            # Essayer de continuer avec la ligne suivante
            continue

    # Sauvegarder le PDF
    try:
        pdf.output(output_path)
        print(f"✅ PDF généré avec succès: {output_path}")
        return True
    except Exception as e:
        print(f"❌ Erreur lors de la sauvegarde: {e}")
        return False


# ================= CONVERSION EN BATCH =================

def batch_convert(folder: str, output_dir: str = "pdf_output") -> int:
    """
    Convertit tous les fichiers DOCX et TXT d'un dossier en PDF
    
    Args:
        folder: Dossier contenant les fichiers à convertir
        output_dir: Dossier de sortie
    
    Returns:
        Nombre de fichiers convertis
    """
    if not os.path.isdir(folder):
        print(f"❌ Dossier introuvable: {folder}")
        return 0

    os.makedirs(output_dir, exist_ok=True)
    
    fichiers_traites = 0
    fichiers_total = 0

    print(f"\n📂 Scan du dossier: {folder}")
    
    for fichier in os.listdir(folder):
        if fichier.lower().endswith((".docx", ".txt")):
            fichiers_total += 1
            print(f"\n{'=' * 70}")
            print(f"Traitement de: {fichier}")
            
            chemin_source = os.path.join(folder, fichier)
            nom_sortie = os.path.splitext(fichier)[0]
            
            if convert_to_pdf(chemin_source, nom_sortie, output_dir):
                fichiers_traites += 1
    
    print(f"\n{'=' * 70}")
    print(f"✅ Conversion terminée: {fichiers_traites}/{fichiers_total} fichiers")
    return fichiers_traites


# ================= FONCTIONS POUR COMPATIBILITÉ AVEC MAIN.PY =================

def convertir_en_pdf(fichier: str):
    """Fonction wrapper pour compatibilité avec main.py"""
    if not check_and_install_dependencies():
        print("❌ Impossible d'installer les dépendances")
        return False
    
    nom_sortie = os.path.splitext(os.path.basename(fichier))[0]
    return convert_to_pdf(fichier, nom_sortie)


# ================= INTERFACE PRINCIPALE =================

def main():
    """Fonction principale - mode standalone"""
    print("\n📄 CYBER FORGE SCAN - Convertisseur PDF")
    print("=" * 70)
    
    # Vérifier et installer les dépendances
    if not check_and_install_dependencies():
        print("\n❌ Installation des dépendances échouée")
        return

    while True:
        print("\n📋 Options:")
        print("1. Convertir un fichier")
        print("2. Convertir un dossier (batch)")
        print("3. À propos")
        print("4. Quitter")

        choix = input("\nVotre choix: ").strip()

        if choix == "1":
            src = input("\n📁 Chemin du fichier (.docx ou .txt): ").strip()
            if not src:
                print("❌ Chemin vide")
                continue
            
            name = input("📝 Nom du PDF (sans extension): ").strip()
            if not name:
                name = os.path.splitext(os.path.basename(src))[0]
            
            out = input("📂 Dossier de sortie (vide = même dossier): ").strip() or None
            
            convert_to_pdf(src, name, out)

        elif choix == "2":
            folder = input("\n📁 Chemin du dossier: ").strip()
            if not folder:
                print("❌ Chemin vide")
                continue
            
            out = input("📂 Dossier de sortie: ").strip() or "pdf_output"
            batch_convert(folder, out)

        elif choix == "3":
            print("\n" + "=" * 70)
            print("📄 CONVERTISSEUR PDF - CYBER FORGE SCAN")
            print("=" * 70)
            print("\n✨ Fonctionnalités:")
            print("  • Conversion DOCX → PDF")
            print("  • Conversion TXT → PDF")
            print("  • Support Unicode (caractères accentués)")
            print("  • Numérotation automatique des pages")
            print("  • Conversion en batch (plusieurs fichiers)")
            print("  • Détection automatique de l'encodage")
            print("\n📦 Dépendances:")
            print("  • fpdf2 (génération PDF)")
            print("  • python-docx (lecture DOCX)")
            print("\n💡 Astuce:")
            print("  Placez DejaVuSans.ttf dans le même dossier")
            print("  pour un meilleur support Unicode")

        elif choix == "4":
            print("\n👋 Au revoir!")
            break

        else:
            print("❌ Choix invalide")


if __name__ == "__main__":
    main()